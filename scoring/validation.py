# -*- coding: utf-8 -*-
"""
ماژول Validation — فاز ۶ (اعتبارسنجی تجربی چارچوب، مورد ۱۲).

کار این ماژول: ترکیب خروجی CSV چارچوب (از reporting/reporter.py) با
نمره‌های انسانی (از data/human_ratings_template.csv پرشده) بر اساس
url، و محاسبه‌ی همبستگی اسپیرمن بین امتیاز ابزار و قضاوت انسانی —
برای هر مؤلفه جداگانه و برای امتیاز نهایی.

⚠️ محدودیت مهم و صادقانه: این تحلیل با فرض «یک داور انسانی» نوشته
شده (نه ۲-۳ داور مستقل که در طرح اولیه پیشنهاد شده بود). به همین
دلیل این ماژول **توافق بین‌داوری (inter-rater agreement / Cohen's
kappa)** را محاسبه نمی‌کند — چون با یک داور، اصلاً همچین مفهومی
تعریف‌شده نیست. اگر در آینده داور دوم/سوم اضافه شد، محاسبه‌ی
Cohen's kappa یا Krippendorff's alpha باید جداگانه اضافه شود؛ فعلاً
عمداً این‌جا نیست تا ادعای اشتباهی درباره‌ی «اعتبارسنجی چند‌داوره»
القا نشود.

مرجع آماری: scipy.stats.spearmanr (🟢 پیاده‌سازی استاندارد و
مستندشده‌ی همبستگی رتبه‌ای اسپیرمن؛ قبلاً هم در requirements.txt این
پروژه بود).
"""

import csv

from scipy.stats import spearmanr


def load_framework_scores(csv_path: str) -> dict:
    """
    خواندن خروجی reporting.export_to_csv و برگرداندن دیکشنری
    {url: {"experience": float|None, "expertise": ..., "authority": ...,
           "trust": ..., "final": ...}}.
    """
    scores = {}
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            url = row["url"]
            scores[url] = {
                "experience": _to_float_or_none(row.get("experience_score")),
                "expertise": _to_float_or_none(row.get("expertise_score")),
                "authority": _to_float_or_none(row.get("authority_score")),
                "trust": _to_float_or_none(row.get("trust_score")),
                "final": _to_float_or_none(row.get("final_score")),
            }
    return scores


def load_human_ratings(csv_path: str) -> dict:
    """
    خواندن data/human_ratings_template.csv پرشده و برگرداندن دیکشنری
    {url: {"experience": int, "expertise": int, "authority": int, "trust": int}}.
    """
    ratings = {}
    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            url = row["url"]
            ratings[url] = {
                "experience": _to_float_or_none(row.get("human_experience")),
                "expertise": _to_float_or_none(row.get("human_expertise")),
                "authority": _to_float_or_none(row.get("human_authority")),
                "trust": _to_float_or_none(row.get("human_trust")),
            }
    return ratings


def _to_float_or_none(value):
    if value is None or value == "":
        return None
    try:
        return float(value)
    except ValueError:
        return None


def compute_spearman_per_component(framework_csv_path: str, human_ratings_csv_path: str) -> dict:
    """
    محاسبه‌ی همبستگی اسپیرمن بین امتیاز چارچوب و نمره‌ی انسانی، برای
    هر ۴ مؤلفه به‌طور جداگانه، بعلاوه یک ترکیب کلی (میانگین ۴ مؤلفه
    انسانی در برابر امتیاز نهایی چارچوب).

    فقط صفحاتی که هم در خروجی چارچوب هم در نمره‌های انسانی، مقدار
    معتبر (نه None) دارند وارد محاسبه می‌شوند. اگر داده‌ی کافی
    (حداقل ۳ جفت) نبود، برای آن مؤلفه None برمی‌گرداند (نه خطا) —
    چون همبستگی روی نمونه‌ی خیلی کوچک بی‌معنی/گمراه‌کننده است.

    خروجی: {"experience": {"rho": float, "p_value": float, "n": int} یا None, ...}
    """
    framework_scores = load_framework_scores(framework_csv_path)
    human_ratings = load_human_ratings(human_ratings_csv_path)

    common_urls = set(framework_scores.keys()) & set(human_ratings.keys())

    results = {}
    for component in ["experience", "expertise", "authority", "trust"]:
        framework_values = []
        human_values = []
        for url in common_urls:
            f_val = framework_scores[url].get(component)
            h_val = human_ratings[url].get(component)
            if f_val is not None and h_val is not None:
                framework_values.append(f_val)
                human_values.append(h_val)

        results[component] = _spearman_or_none(framework_values, human_values)

    # ترکیب کلی: میانگین ۴ نمره‌ی انسانی در برابر امتیاز نهایی چارچوب
    final_framework_values = []
    overall_human_values = []
    for url in common_urls:
        f_final = framework_scores[url].get("final")
        h_components = [human_ratings[url].get(c) for c in
                         ["experience", "expertise", "authority", "trust"]]
        if f_final is not None and all(v is not None for v in h_components):
            final_framework_values.append(f_final)
            overall_human_values.append(sum(h_components) / len(h_components))

    results["overall"] = _spearman_or_none(final_framework_values, overall_human_values)
    return results


def _spearman_or_none(x_values: list, y_values: list, min_pairs: int = 3):
    if len(x_values) < min_pairs or len(y_values) < min_pairs:
        return None
    rho, p_value = spearmanr(x_values, y_values)
    return {"rho": rho, "p_value": p_value, "n": len(x_values)}


def print_validation_report(results: dict) -> None:
    """چاپ خوانا از خروجی compute_spearman_per_component."""
    component_labels = {
        "experience": "Experience",
        "expertise": "Expertise",
        "authority": "Authoritativeness",
        "trust": "Trustworthiness",
        "overall": "امتیاز نهایی (میانگین انسانی در برابر final_score)",
    }
    print("=== گزارش اعتبارسنجی: همبستگی اسپیرمن (ابزار در برابر داور انسانی) ===")
    print("⚠️ این نتایج فقط بر پایه‌ی یک داور انسانی است — نه چند داور مستقل.")
    print()
    for component, label in component_labels.items():
        result = results.get(component)
        if result is None:
            print(f"{label}: داده‌ی کافی نبود (حداقل ۳ صفحه‌ی مشترک با نمره‌ی معتبر لازم است)")
        else:
            print(f"{label}: rho={result['rho']:.3f}, p={result['p_value']:.4f}, n={result['n']}")


# ===========================================================================
# تحلیل درون‌حافظه‌ای برای صفحه‌ی «اعتبارسنجی» وب‌اپ (webapp/validation_web.py)
# همان آمار بالا، ولی روی متن فایل‌های بارگذاری‌شده، با فاصله‌ی اطمینان و نمودار.
# ===========================================================================

import io
import math
from urllib.parse import unquote

from scipy.stats import kendalltau, norm

COMPONENT_KEYS = ["experience", "expertise", "authority", "trust"]
HUMAN_COLUMNS = {
    "experience": "human_experience",
    "expertise": "human_expertise",
    "authority": "human_authority",
    "trust": "human_trust",
}
HUMAN_MIN, HUMAN_MAX = 1, 5
MIN_PAIRS_FOR_CI = 10      # روش فاصله‌ی اطمینان زیر این تعداد قابل‌اتکا نیست
SMALL_SAMPLE = 30          # زیر این تعداد، نتیجه فقط اولیه است
XLSX_MAX_ROWS = 20000
XLSX_MAX_COLS = 50


class ValidationInputError(ValueError):
    """خطای قابل‌نمایش به کاربر (فایل ناخوانا، ستون‌های لازم نبود، ...)."""


def _cell_text(value) -> str:
    """مقدار یک خانه‌ی اکسل را به متن CSV تبدیل می‌کند (۳٫۰ → «3»)."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def xlsx_bytes_to_csv_text(raw: bytes) -> str:
    """
    برگه‌ی اول یک فایل .xlsx را به متن CSV تبدیل می‌کند، تا بقیه‌ی منطق
    (خواندن ستون‌ها، مطابقت آدرس‌ها، اعتبارسنجی نمره‌ها) برای هر دو قالب یکی باشد.
    سقف ردیف و ستون برای جلوگیری از فایل‌های عجیب.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as e:  # کتابخانه‌ی خواندن اکسل روی این رایانه نصب نیست
        raise ValidationInputError(
            "کتابخانه‌ی openpyxl روی این رایانه نصب نیست. در ترمینال بنویسید: "
            "pip install openpyxl ؛ سپس برنامه را دوباره اجرا کنید. (یا فایل را به‌صورت CSV UTF-8 بارگذاری کنید.)"
        ) from e
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        out = io.StringIO()
        writer = csv.writer(out, lineterminator="\r\n")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= XLSX_MAX_ROWS:
                raise ValidationInputError(f"فایل اکسل بیش از {XLSX_MAX_ROWS} ردیف دارد.")
            writer.writerow([_cell_text(v) for v in row[:XLSX_MAX_COLS]])
        wb.close()
        return out.getvalue()
    except ValidationInputError:
        raise
    except Exception as e:  # فایل خراب یا غیرِ اکسل
        raise ValidationInputError("فایل اکسل خوانده نشد؛ شاید خراب است یا قالبش درست نیست.") from e


def decode_csv_bytes(raw: bytes) -> str:
    """متن فایل را می‌خواند؛ BOM اکسل را هم می‌پذیرد و فایل .xlsx را هم قبول می‌کند."""
    if raw[:2] == b"PK":  # فایل .xlsx یک بسته‌ی zip است
        return xlsx_bytes_to_csv_text(raw)
    try:
        return raw.decode("utf-8-sig")
    except UnicodeDecodeError as e:
        raise ValidationInputError(
            "فایل با رمزنگاری UTF-8 ذخیره نشده است. در اکسل هنگام ذخیره «CSV UTF-8» را انتخاب کنید."
        ) from e


def normalize_url_key(url: str) -> str:
    """کلید یکسان برای مطابقت آدرس‌ها: بدون اسکیم، بدون www، بدون / پایانی، حروف کوچک."""
    u = unquote((url or "").strip()).lower()  # %D8%A7 و «ا» یکی حساب شوند
    for prefix in ("https://", "http://"):
        if u.startswith(prefix):
            u = u[len(prefix):]
    if u.startswith("www."):
        u = u[4:]
    return u.rstrip("/")


def _read_rows(text: str, required: list, what: str) -> list:
    reader = csv.DictReader(io.StringIO(text))
    fields = [f.strip() for f in (reader.fieldnames or [])]
    missing = [c for c in required if c not in fields]
    if missing:
        raise ValidationInputError(
            f"در فایل {what} این ستون‌ها پیدا نشد: {'، '.join(missing)}. "
            f"ستون‌های موجود: {'، '.join(fields) or 'هیچ'}."
        )
    rows = []
    for row in reader:
        rows.append({(k or "").strip(): (v or "").strip() for k, v in row.items()})
    return rows


def _unique_urls(framework_csv_text: str) -> list:
    rows = _read_rows(framework_csv_text, ["url"], "نتایج ابزار")
    seen, urls = set(), []
    for r in rows:
        key = normalize_url_key(r["url"])
        if r["url"] and key not in seen:
            seen.add(key)
            urls.append(r["url"])
    if not urls:
        raise ValidationInputError("هیچ آدرسی در فایل نتایج ابزار پیدا نشد.")
    return urls


def build_blank_rating_sheet(framework_csv_text: str) -> str:
    """
    برگه‌ی نمره‌دهی کور: فقط آدرس‌ها و ستون‌های خالی نمره‌ی انسانی — بدون هیچ
    امتیازی از ابزار، تا داور تحت تأثیر نمره‌ی ابزار قرار نگیرد.
    """
    urls = _unique_urls(framework_csv_text)
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\r\n")
    writer.writerow(["url"] + list(HUMAN_COLUMNS.values()))
    for u in urls:
        writer.writerow([u, "", "", "", ""])
    return out.getvalue()


# عرض ستون‌ها به نویسه؛ کاربر می‌تواند در اکسل با کشیدن لبه‌ی ستون تغییرش دهد.
XLSX_URL_WIDTH = 46
XLSX_SCORE_WIDTH = 20
XLSX_ROW_HEIGHT = 20
_XLSX_HEADER_HINTS = {
    "url": "آدرس صفحه. روی آن بزنید تا صفحه باز شود. این ستون را تغییر ندهید.",
    "human_experience": "تجربه: عدد صحیح ۱ تا ۵",
    "human_expertise": "تخصص: عدد صحیح ۱ تا ۵",
    "human_authority": "اقتدار: عدد صحیح ۱ تا ۵",
    "human_trust": "اعتماد: عدد صحیح ۱ تا ۵",
}
_XLSX_GUIDE = [
    "راهنمای نمره‌دهی",
    "۱. برای هر صفحه، در چهار ستون سمت راست آدرس، یک عدد صحیح از ۱ تا ۵ بنویسید (۱ خیلی ضعیف، ۵ خیلی خوب).",
    "۲. معیار هر عدد در فایل data/rating_rubric.md پروژه آمده است؛ پیش از شروع بخوانیدش.",
    "۳. به نمره‌ی ابزار نگاه نکنید. هر صفحه را فقط یک بار نمره بدهید و بعداً عوض نکنید.",
    "۴. عدد اعشاری و خانه‌ی خالی نگذارید. اگر عدد نادرست بنویسید، اکسل خطا می‌دهد.",
    "۵. نام ستون‌های ردیف اول را تغییر ندهید، وگرنه برنامه فایل را نمی‌شناسد.",
    "۶. عرض ستون‌ها را هر طور راحت‌اید با کشیدن لبه‌ی ستون تغییر دهید.",
    "۷. پس از پایان، همین فایل را با همین قالب (.xlsx) ذخیره کنید و در فرم دوم صفحه‌ی اعتبارسنجی بارگذاری کنید.",
]


def build_blank_rating_xlsx(framework_csv_text: str) -> bytes:
    """
    همان برگه‌ی نمره‌دهی کور، اما به‌صورت فایل اکسل: عرض ستون‌ها ثابت و قابل‌تغییر،
    آدرس‌ها بریده‌شده در دو خط (نه سرریز به ستون کناری) و قابل‌کلیک، ردیف عنوان ثابت،
    و فقط عدد صحیح ۱ تا ۵ در خانه‌های نمره پذیرفته می‌شود.
    فایل CSV نمی‌تواند عرض ستون ذخیره کند؛ برای همین این قالب لازم است.
    """
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    urls = _unique_urls(framework_csv_text)
    headers = ["url"] + list(HUMAN_COLUMNS.values())

    wb = Workbook()
    ws = wb.active
    ws.title = "نمره‌دهی"

    head_fill = PatternFill("solid", fgColor="34443D")
    head_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="DCDDD5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for col, name in enumerate(headers, start=1):
        c = ws.cell(row=1, column=col, value=name)
        c.fill, c.font, c.border = head_fill, head_font, border
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.comment = Comment(_XLSX_HEADER_HINTS[name], "E-E-A-T")
        ws.column_dimensions[get_column_letter(col)].width = (
            XLSX_URL_WIDTH if col == 1 else XLSX_SCORE_WIDTH)
    ws.row_dimensions[1].height = 24

    link_font = Font(color="0563C1", underline="single")
    for r, u in enumerate(urls, start=2):
        # آدرس اصلی (فقط حروف لاتین)؛ آدرس فارسی‌شده در یک خانه راست‌به‌چپ و چپ‌به‌راست
        # را درهم می‌ریزد. با کلیک روی خانه صفحه در مرورگر باز می‌شود.
        c = ws.cell(row=r, column=1, value=u)
        if len(u) <= 2000:  # سقف طول پیوند در اکسل
            c.hyperlink = u
        c.font = link_font
        # شکستن خط + ارتفاع ثابت: متن بلند بریده می‌شود و روی ستون کناری نمی‌ریزد
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
        c.border = border
        ws.row_dimensions[r].height = XLSX_ROW_HEIGHT
        for col in range(2, 6):
            sc = ws.cell(row=r, column=col)
            sc.alignment = Alignment(horizontal="center", vertical="center")
            sc.border = border

    last = len(urls) + 1
    dv = DataValidation(
        type="whole", operator="between", formula1=str(HUMAN_MIN), formula2=str(HUMAN_MAX),
        allow_blank=True, showErrorMessage=True,
        errorTitle="عدد نادرست", error="فقط عدد صحیح از ۱ تا ۵ مجاز است.",
        showInputMessage=True, promptTitle="نمره", prompt="عدد صحیح ۱ تا ۵",
    )
    dv.add(f"B2:E{last}")
    ws.add_data_validation(dv)
    ws.freeze_panes = "B2"

    guide = wb.create_sheet("راهنما")
    guide.sheet_view.rightToLeft = True
    guide.column_dimensions["A"].width = 110
    for i, line in enumerate(_XLSX_GUIDE, start=1):
        c = guide.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical="top", readingOrder=2)
        if i == 1:
            c.font = Font(bold=True, size=13)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def strength_label(rho: float) -> str:
    """
    برچسب فارسی قدرت همبستگی بر پایه‌ی قاعده‌ی قراردادی رایج (Evans, 1996):
    <0.20 بسیار ضعیف، <0.40 ضعیف، <0.60 متوسط، <0.80 قوی، بقیه بسیار قوی.
    این مرزها قرارداد هستند، نه قانون آماری؛ در متن پایان‌نامه با همین قید بیاید.
    """
    a = abs(rho)
    if a < 0.20:
        word = "بسیار ضعیف"
    elif a < 0.40:
        word = "ضعیف"
    elif a < 0.60:
        word = "متوسط"
    elif a < 0.80:
        word = "قوی"
    else:
        word = "بسیار قوی"
    return word if rho >= 0 else f"{word} (منفی)"


def analyze_pairs(tool_values: list, human_values: list) -> dict | None:
    """
    اسپیرمن (و تاو-بی کندال به‌عنوان مکمل، چون نمره‌ی انسانی ۱ تا ۵ پر از مقدار
    تکراری است) + فاصله‌ی اطمینان ۹۵٪ تقریبی (Bonett & Wright, 2000) وقتی n>=10.
    اگر کمتر از ۳ جفت یا واریانس صفر (همه‌ی مقدارها برابر) باشد None برمی‌گرداند.
    """
    n = len(tool_values)
    if n < 3 or len(set(tool_values)) < 2 or len(set(human_values)) < 2:
        return None
    rho, p = spearmanr(tool_values, human_values)
    if rho != rho:  # NaN
        return None
    tau, tau_p = kendalltau(tool_values, human_values)
    ci = None
    if n >= MIN_PAIRS_FOR_CI and abs(rho) < 1:
        se = math.sqrt((1 + rho ** 2 / 2) / (n - 3))
        z = math.atanh(rho)
        zc = norm.ppf(0.975)
        ci = (math.tanh(z - zc * se), math.tanh(z + zc * se))
    return {"rho": float(rho), "p_value": float(p), "n": n,
            "ci_low": ci[0] if ci else None, "ci_high": ci[1] if ci else None,
            "tau": float(tau), "tau_p": float(tau_p),
            "strength": strength_label(float(rho)),
            # نقطه‌های نمودار: (نمره‌ی انسانی، نمره‌ی ابزار در بازه‌ی ۰ تا ۱)
            "points": [(h, t / (100.0 if max(tool_values) > 1.0 else 1.0))
                       for h, t in zip(human_values, tool_values)]}


def validate_uploaded(framework_text: str, human_text: str) -> dict:
    """
    ورودی: متن CSV خروجی ابزار (ستون‌های url و *_score) و متن برگه‌ی نمره‌ی انسانی.
    خروجی: نتیجه‌ی هر مؤلفه و امتیاز نهایی + شمارش صفحه‌های کنار رفته (با دلیل) +
    هشدارهای صادقانه. نمره‌ی ابزار از ۰ تا ۱ و نمره‌ی انسانی ۱ تا ۵ است؛ اسپیرمن
    فقط به ترتیب رتبه‌ها نگاه می‌کند، پس مقیاس‌ها لازم نیست یکی شوند.
    """
    tool_cols = {"experience": "experience_score", "expertise": "expertise_score",
                 "authority": "authority_score", "trust": "trust_score", "final": "final_score"}
    fw_rows = _read_rows(framework_text, ["url"] + list(tool_cols.values()), "نتایج ابزار")
    hu_rows = _read_rows(human_text, ["url"] + list(HUMAN_COLUMNS.values()), "نمره‌ی داور")

    tool = {}
    for r in fw_rows:
        tool[normalize_url_key(r["url"])] = {k: _to_float_or_none(r.get(c)) for k, c in tool_cols.items()}

    human, bad_values, duplicates = {}, 0, 0
    for r in hu_rows:
        key = normalize_url_key(r["url"])
        if not key:
            continue
        if key in human:
            duplicates += 1
            continue
        vals = {}
        for comp, col in HUMAN_COLUMNS.items():
            v = _to_float_or_none(r.get(col))
            if v is not None and not (HUMAN_MIN <= v <= HUMAN_MAX):
                bad_values += 1
                v = None
            vals[comp] = v
        human[key] = vals

    common = [k for k in human if k in tool]
    rated = [k for k in human if any(v is not None for v in human[k].values())]
    only_human = [k for k in rated if k not in tool]

    results = {}
    for comp in COMPONENT_KEYS:
        xs, ys = [], []
        for k in common:
            t, h = tool[k][comp], human[k][comp]
            if t is not None and h is not None:
                xs.append(t)
                ys.append(h)
        results[comp] = analyze_pairs(xs, ys)

    xs, ys = [], []
    for k in common:
        hs = [human[k][c] for c in COMPONENT_KEYS]
        if tool[k]["final"] is not None and all(v is not None for v in hs):
            xs.append(tool[k]["final"])
            ys.append(sum(hs) / len(hs))
    results["overall"] = analyze_pairs(xs, ys)

    n_overall = results["overall"]["n"] if results["overall"] else 0
    warnings = ["نمره‌ها فقط از یک داور است؛ توافق بین‌داوری قابل محاسبه نیست و نتیجه نشان می‌دهد ابزار "
                "با قضاوت همان یک نفر چقدر هم‌جهت است، نه با یک اجماع."]
    if n_overall < SMALL_SAMPLE:
        warnings.append(f"تعداد صفحه‌های کامل ({n_overall}) از {SMALL_SAMPLE} کمتر است؛ نتیجه فقط اولیه است "
                        f"و فاصله‌ی اطمینان پهن می‌ماند.")
    warnings.append("پنج آزمون هم‌زمان (چهار مؤلفه و امتیاز نهایی) انجام شده است؛ مقدار احتمال هر کدام بدون "
                    "تصحیح چندآزمونی است و نباید تک‌تک و خوش‌بینانه تفسیر شود.")
    warnings.append("نمره‌ی انسانی عدد صحیح ۱ تا ۵ است و مقدارهای تکراری زیاد دارد؛ برای همین تاو-بی کندال هم "
                    "کنار اسپیرمن آمده است.")
    if bad_values:
        warnings.append(f"{bad_values} نمره خارج از بازه‌ی ۱ تا ۵ بود و کنار گذاشته شد.")
    if duplicates:
        warnings.append(f"{duplicates} ردیف تکراری در برگه‌ی نمره‌دهی بود؛ فقط اولین ردیف هر آدرس شمرده شد.")

    return {
        "results": results,
        "counts": {"tool_pages": len(tool), "rated_pages": len(rated), "matched": len(common),
                   "only_human": len(only_human), "unrated": len(human) - len(rated)},
        "warnings": warnings,
    }


_FA_DIGITS = str.maketrans("0123456789.", "۰۱۲۳۴۵۶۷۸۹٫")


def _fa(text) -> str:
    """عدد را با رقم و ممیز فارسی می‌نویسد (متن فارسی بدون قاطی‌شدن جهت نوشتار)."""
    return str(text).translate(_FA_DIGITS)


def thesis_paragraph(report: dict) -> str:
    """پاراگراف پیش‌نویس برای فصل نتایج (فقط از عددهای واقعی محاسبه‌شده؛ چیزی ساخته نمی‌شود)."""
    labels = {"experience": "تجربه", "expertise": "تخصص", "authority": "اقتدار",
              "trust": "اعتماد", "overall": "امتیاز نهایی"}
    n = report["counts"]["matched"]
    parts = [f"برای سنجش اعتبار، امتیاز ابزار با نمره‌ی یک داور انسانی روی {_fa(n)} صفحه مقایسه شد "
             f"(نمره‌ی انسانی ۱ تا ۵، همبستگی رتبه‌ای اسپیرمن)."]
    for key in ("experience", "expertise", "authority", "trust", "overall"):
        r = report["results"].get(key)
        if r is None:
            parts.append(f"برای {labels[key]} داده‌ی کافی نبود.")
            continue
        ci = (f"، فاصله‌ی اطمینان ۹۵٪ تقریبی از {_fa(format(r['ci_low'], '.2f'))} "
              f"تا {_fa(format(r['ci_high'], '.2f'))}" if r["ci_low"] is not None else "")
        p_txt = ("کمتر از ۰٫۰۰۱" if r["p_value"] < 0.001 else _fa(format(r["p_value"], ".3f")))
        parts.append(f"{labels[key]}: ضریب اسپیرمن {_fa(format(r['rho'], '.2f'))} ({r['strength']})، "
                     f"روی {_fa(r['n'])} صفحه، مقدار احتمال {p_txt}{ci}.")
    parts.append("چون نمره‌دهی تنها یک داور داشته، این نتیجه توافق با همان داور را نشان می‌دهد و "
                 "اجماع چندنفره نیست؛ مقدارهای احتمال تصحیح چندآزمونی ندارند.")
    return " ".join(parts)
