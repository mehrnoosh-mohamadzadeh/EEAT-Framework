"""
تست خواندن فایل‌های CSV که اکسل با «CSV UTF-8» ذخیره می‌کند
(اکسل یک نشانه‌ی نامرئی BOM به ابتدای فایل می‌افزاید).
"""

import csv

import pytest

from main import read_urls_from_csv
from scoring.validation import load_framework_scores, load_human_ratings


def _write(path, header, rows, bom):
    enc = "utf-8-sig" if bom else "utf-8"
    with open(path, "w", encoding=enc, newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


def test_main_reads_url_column_with_and_without_bom(tmp_path):
    for bom in (True, False):
        p = tmp_path / f"urls_{bom}.csv"
        _write(p, ["url", "category"], [["https://example.com/a", "x"]], bom)
        rows = read_urls_from_csv(str(p))
        assert rows[0]["url"] == "https://example.com/a"


def test_validation_loaders_accept_bom(tmp_path):
    p1 = tmp_path / "res.csv"
    _write(p1, ["url", "experience_score", "expertise_score", "authority_score",
                "trust_score", "final_score"], [["https://e.com/1", "50", "60", "70", "80", "65"]], True)
    p2 = tmp_path / "hum.csv"
    _write(p2, ["url", "human_experience", "human_expertise", "human_authority",
                "human_trust"], [["https://e.com/1", "3", "4", "5", "4"]], True)
    assert load_framework_scores(str(p1))["https://e.com/1"]["final"] == 65.0
    assert load_human_ratings(str(p2))["https://e.com/1"]["authority"] == 5.0


def test_shipped_urls_file_is_readable_and_unique():
    rows = read_urls_from_csv("data/urls.csv")
    urls = [r["url"] for r in rows]
    assert len(urls) >= 30
    assert len(set(urls)) == len(urls)
    assert all(u.startswith("https://") for u in urls)


# ---------- برگه‌ی اکسل نمره‌دهی: آدرس خوانا و مطابقت با آدرس درصدی ----------

import io

from openpyxl import load_workbook

from scoring.validation import build_blank_rating_xlsx, normalize_url_key

_ENC = "https://site.test/%D8%AE%D8%A8%D8%B1-%D8%AA%D8%B3%D8%AA"
_DEC = "https://site.test/خبر-تست"


def test_url_key_same_for_encoded_and_decoded():
    assert normalize_url_key(_ENC) == normalize_url_key(_DEC)


def test_blank_xlsx_keeps_original_url_and_links_to_it():
    text = "url,final_score\r\n" + _ENC + ",50\r\n"
    wb = load_workbook(io.BytesIO(build_blank_rating_xlsx(text)))
    ws = wb.worksheets[0]
    assert ws["A2"].value == _ENC
    assert ws["A2"].hyperlink.target == _ENC
    # عرض ستون‌ها ثابت است و ردیف ارتفاع ثابت دارد (متن بلند بریده می‌شود)
    assert ws.column_dimensions["A"].width and ws.column_dimensions["B"].width
    assert ws["A2"].alignment.wrap_text is True
    assert ws.row_dimensions[2].height


def test_filled_xlsx_matches_results_with_encoded_urls():
    from scoring.validation import validate_uploaded, decode_csv_bytes
    header = "url,category,experience_score,expertise_score,authority_score,trust_score,final_score\r\n"
    rows = [(_ENC + str(i), 10 * i) for i in range(1, 6)]
    fw = header + "".join(f"{u},x,{v},{v},{v},{v},{v}\r\n" for u, v in rows)
    wb = load_workbook(io.BytesIO(build_blank_rating_xlsx(fw)))
    ws = wb.worksheets[0]
    for r in range(2, 7):
        for c in range(2, 6):
            ws.cell(row=r, column=c, value=r - 1)
    buf = io.BytesIO()
    wb.save(buf)
    report = validate_uploaded(fw, decode_csv_bytes(buf.getvalue()))
    assert report is not None


def test_missing_openpyxl_gives_clear_message(monkeypatch):
    """اگر openpyxl نصب نباشد، پیام باید علت واقعی را بگوید، نه «فایل خراب است»."""
    import builtins
    from scoring.validation import ValidationInputError, xlsx_bytes_to_csv_text
    real_import = builtins.__import__

    def fake_import(name, *a, **k):
        if name.startswith("openpyxl"):
            raise ImportError("No module named openpyxl")
        return real_import(name, *a, **k)

    monkeypatch.setattr(builtins, "__import__", fake_import)
    with pytest.raises(ValidationInputError) as ei:
        xlsx_bytes_to_csv_text(b"PK\x03\x04")
    assert "pip install openpyxl" in str(ei.value)
